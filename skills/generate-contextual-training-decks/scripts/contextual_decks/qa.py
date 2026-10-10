from __future__ import annotations

import json
import math
import shutil
import subprocess
from dataclasses import asdict, dataclass, field
from pathlib import Path

from pptx import Presentation

from .models import DeckSpec


@dataclass
class QAItem:
    level: str
    code: str
    message: str
    slide: int | None = None


@dataclass
class QAReport:
    passed: bool = True
    slide_count: int = 0
    layout_sequence: list[str] = field(default_factory=list)
    items: list[QAItem] = field(default_factory=list)
    rendered_pdf: str = ""

    def add(self, level: str, code: str, message: str, slide: int | None = None) -> None:
        self.items.append(QAItem(level=level, code=code, message=message, slide=slide))
        if level == "error":
            self.passed = False

    def to_dict(self) -> dict:
        return asdict(self)


def _luminance(hex_color: str) -> float:
    values = [int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    linear = [v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in values]
    return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]


def contrast_ratio(a: str, b: str) -> float:
    light, dark = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (light + 0.05) / (dark + 0.05)


def inspect_spec(deck: DeckSpec, layout_sequence: list[str] | None = None) -> QAReport:
    report = QAReport(slide_count=len(deck.slides), layout_sequence=layout_sequence or [])
    if len(deck.slides) < 3:
        report.add("warning", "deck.too_short", "The deck contains fewer than three slides.")

    for index, slide in enumerate(deck.slides, start=1):
        title_words = len(slide.title.split())
        if title_words > 16:
            report.add("warning", "title.long", f"Title has {title_words} words.", index)
        visible_words = len(" ".join([slide.body, *slide.bullets]).split())
        visible_words += sum(len(" ".join([item.title, item.text, item.note]).split()) for item in slide.items)
        if visible_words > 110:
            report.add("warning", "content.dense", f"Visible content has {visible_words} words.", index)
        if slide.kind.value == "activity" and not slide.notes.instruction:
            report.add("error", "activity.no_instruction", "Activity has no facilitator instruction.", index)
        if slide.kind.value == "quiz" and (not slide.choices or slide.answer_index is None):
            report.add("error", "quiz.incomplete", "Quiz must include choices and answer_index.", index)

    if layout_sequence:
        run_length = 1
        for index in range(1, len(layout_sequence)):
            if layout_sequence[index] == layout_sequence[index - 1]:
                run_length += 1
                if run_length >= 3:
                    report.add(
                        "warning",
                        "layout.repeated",
                        f"Layout '{layout_sequence[index]}' repeats {run_length} times.",
                        index + 1,
                    )
            else:
                run_length = 1

    themes = {"default": deck.theme, **deck.theme_variants}
    for theme_key, theme in themes.items():
        for foreground, background, label in (
            (theme.ink, theme.background, "ink/background"),
            (theme.paper, theme.primary, "paper/primary"),
        ):
            ratio = contrast_ratio(foreground, background)
            if ratio < 4.5:
                report.add(
                    "error",
                    "contrast.low",
                    f"Theme '{theme_key}' contrast {label} is {ratio:.2f}:1.",
                )
    return report


def inspect_pptx(path: str | Path, report: QAReport) -> QAReport:
    presentation = Presentation(str(path))
    if len(presentation.slides) != report.slide_count:
        report.add(
            "error",
            "pptx.slide_count_mismatch",
            f"PPTX has {len(presentation.slides)} slides; expected {report.slide_count}.",
        )
    for index, slide in enumerate(presentation.slides, start=1):
        text = " ".join(
            shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        )
        if not text:
            report.add("error", "pptx.empty_slide", "Rendered slide has no text.", index)
    return report


def render_pdf(pptx_path: str | Path, output_dir: str | Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output),
            str(Path(pptx_path).resolve()),
        ],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    expected = output / (Path(pptx_path).stem + ".pdf")
    return expected if completed.returncode == 0 and expected.exists() else None


def write_report(report: QAReport, path: str | Path) -> None:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
