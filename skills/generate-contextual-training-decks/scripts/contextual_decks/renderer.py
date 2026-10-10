from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .models import ContentItem, DeckSpec, SlideKind, SlideSpec, ThemeSpec


SLIDE_W = 13.333
SLIDE_H = 7.5
SAFE_X = 0.58
CONTENT_TOP = 1.48
CONTENT_BOTTOM = 6.92


def _rgb(value: str) -> RGBColor:
    cleaned = value.lstrip("#")
    return RGBColor.from_string(cleaned)


def _is_arabic(text: str) -> bool:
    return bool(re.search(r"[\u0600-\u06FF]", text))


def _set_run_font(run, name: str, size: float, color: str, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    r_pr = run._r.get_or_add_rPr()
    r_pr.set("lang", "ar-SA" if _is_arabic(run.text) else "en-US")
    for tag in ("a:latin", "a:ea", "a:cs"):
        child = r_pr.find(qn(tag))
        if child is None:
            child = r_pr.makeelement(qn(tag))
            r_pr.append(child)
        child.set("typeface", name)


def _set_paragraph_direction(paragraph, rtl: bool) -> None:
    paragraph.alignment = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("rtl", "1" if rtl else "0")


def _set_alt_text(shape, text: str) -> None:
    if not text:
        return
    try:
        c_nv_pr = shape._element.xpath(".//p:cNvPr")[0]
        c_nv_pr.set("descr", text[:1000])
    except (IndexError, AttributeError):
        return


def _solid_fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)


def _line(shape, color: str, width: float = 1.0) -> None:
    shape.line.color.rgb = _rgb(color)
    shape.line.width = Pt(width)


def _no_line(shape) -> None:
    shape.line.fill.background()


def _add_shape(slide, shape_type, x, y, w, h, fill: str, line: str | None = None, radius=True):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    _solid_fill(shape, fill)
    _line(shape, line or fill, 0.8) if line else _no_line(shape)
    return shape


def _add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    theme: ThemeSpec,
    size: float = 24,
    color: str | None = None,
    bold: bool = False,
    align: PP_ALIGN | None = None,
    rtl: bool | None = None,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.04,
    name: str = "",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    direction = _is_arabic(text) if rtl is None else rtl
    _set_paragraph_direction(paragraph, direction)
    if align is not None:
        paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.space_before = Pt(0)
    paragraph.line_spacing = 1.05
    if paragraph.runs:
        _set_run_font(
            paragraph.runs[0],
            theme.font_bold if bold else theme.font_regular,
            size,
            color or theme.ink,
            bold,
        )
    _set_alt_text(box, text)
    return box


def _add_multiline(
    slide,
    lines: Iterable[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    theme: ThemeSpec,
    size: float = 22,
    color: str | None = None,
    bullet: bool = False,
    numbered: bool = False,
    spacing: float = 9,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        prefix = f"{index + 1}.  " if numbered else ("•  " if bullet else "")
        paragraph.text = prefix + line
        _set_paragraph_direction(paragraph, _is_arabic(line))
        paragraph.space_after = Pt(spacing)
        paragraph.line_spacing = 1.08
        if paragraph.runs:
            _set_run_font(paragraph.runs[0], theme.font_regular, size, color or theme.ink)
    _set_alt_text(box, " | ".join(lines))
    return box


def _add_circle_label(slide, label: str, x: float, y: float, d: float, theme: ThemeSpec, *, fill=None, color=None, size=20):
    circle = _add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill or theme.primary)
    _add_text(
        slide,
        label,
        x,
        y + 0.01,
        d,
        d - 0.02,
        theme=theme,
        size=size,
        color=color or theme.paper,
        bold=True,
        align=PP_ALIGN.CENTER,
        rtl=_is_arabic(label),
        valign=MSO_ANCHOR.MIDDLE,
    )
    return circle


def _item_text(item: ContentItem) -> str:
    return item.text or item.note or item.title or item.label or item.value


def _fit_items(items: list[ContentItem], fallback: list[str], minimum: int = 2) -> list[ContentItem]:
    if items:
        return items
    generated = [ContentItem(title=text) for text in fallback]
    return generated if generated else [ContentItem(title="فكرة رئيسية") for _ in range(minimum)]


def _slide_background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _add_header(slide, spec: SlideSpec, theme: ThemeSpec, number: int, *, light=False) -> None:
    ink = theme.paper if light else theme.ink
    muted = theme.paper if light else theme.muted
    eyebrow = spec.eyebrow or spec.kind.value.replace("_", " ")
    _add_text(slide, eyebrow, 8.4, 0.47, 4.25, 0.35, theme=theme, size=13, color=muted, bold=True)
    title_size = 35 if len(spec.title) <= 58 else 30
    _add_text(slide, spec.title, 0.85, 0.76, 11.8, 0.68, theme=theme, size=title_size, color=ink, bold=True)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.68), Inches(0.76), Inches(0.08), Inches(0.55))
    _solid_fill(bar, theme.accent)
    _no_line(bar)


def _add_footer(
    slide, deck: DeckSpec, number: int, *, light=False, theme: ThemeSpec | None = None
) -> None:
    theme = theme or deck.theme
    color = theme.paper if light else theme.muted
    label = theme.footer_text or deck.program_name or deck.title
    _add_text(slide, label, 7.65, 7.03, 5.05, 0.24, theme=theme, size=10.5, color=color)
    _add_text(
        slide,
        f"{number:02d}",
        0.58,
        7.02,
        0.55,
        0.24,
        theme=theme,
        size=10.5,
        color=color,
        bold=True,
        align=PP_ALIGN.LEFT,
        rtl=False,
    )


def _add_notes(slide, spec: SlideSpec) -> None:
    notes = spec.notes
    parts = [
        f"الهدف | {notes.objective}",
        f"الزمن | {notes.timing}",
        f"ما يقوله الميسّر | {notes.script}",
        f"سؤال/تعليمات | {notes.instruction}",
        f"الإجابة أو ملاحظة التيسير | {notes.answer_or_debrief}",
        f"الانتقال | {notes.transition}",
    ]
    if spec.sources:
        parts.append("[Sources]")
        for source in spec.sources:
            suffix = f" | {source.accessed}" if source.accessed else ""
            parts.append(f"- {source.title} | {source.url}{suffix}")
    frame = slide.notes_slide.notes_text_frame
    if frame is not None:
        frame.text = "\n".join(parts)


class LayoutSelector:
    variants = {
        SlideKind.COVER: ["cover-hero", "cover-split"],
        SlideKind.SECTION: ["section-number", "section-band"],
        SlideKind.STATEMENT: ["statement-center", "statement-sidebar"],
        SlideKind.CONCEPT: ["concept-cards", "concept-split", "concept-focus"],
        SlideKind.PROCESS: ["process-horizontal", "process-stair"],
        SlideKind.TIMELINE: ["timeline-ribbon", "timeline-lanes"],
        SlideKind.COMPARISON: ["comparison-columns", "comparison-balance"],
        SlideKind.METRICS: ["metrics-grid", "metrics-hero"],
        SlideKind.CASE_STUDY: ["case-decision", "case-scenario"],
        SlideKind.ACTIVITY: ["activity-workshop", "activity-canvas"],
        SlideKind.QUIZ: ["quiz-options"],
        SlideKind.TOOL: ["tool-canvas", "tool-steps"],
        SlideKind.SUMMARY: ["summary-checklist", "summary-commitment"],
    }

    def __init__(self) -> None:
        self.previous = ""
        self.usage: dict[str, int] = {}

    def infer_kind(self, spec: SlideSpec, index: int) -> SlideKind:
        if spec.kind != SlideKind.AUTO:
            return spec.kind
        text = f"{spec.title} {spec.subtitle} {spec.body}".lower()
        if index == 1:
            return SlideKind.COVER
        if spec.choices:
            return SlideKind.QUIZ
        if len(spec.items) >= 3 and any(item.value for item in spec.items):
            return SlideKind.METRICS
        if any(token in text for token in ("خطوات", "عملية", "workflow", "process")):
            return SlideKind.PROCESS
        if any(token in text for token in ("مقارنة", "مقابل", "versus", " vs ")):
            return SlideKind.COMPARISON
        if any(token in text for token in ("نشاط", "تطبيق", "ورشة", "activity")):
            return SlideKind.ACTIVITY
        if any(token in text for token in ("حالة", "سيناريو", "case")):
            return SlideKind.CASE_STUDY
        if len(spec.items) >= 2:
            return SlideKind.CONCEPT
        if len(spec.body.split()) <= 22 and not spec.bullets:
            return SlideKind.STATEMENT
        return SlideKind.CONCEPT

    def select(self, spec: SlideSpec, index: int) -> tuple[SlideKind, str]:
        kind = self.infer_kind(spec, index)
        options = self.variants[kind]
        if spec.layout_hint in options:
            chosen = spec.layout_hint
        else:
            ranked = sorted(options, key=lambda name: (name == self.previous, self.usage.get(name, 0), name))
            chosen = ranked[0]
        self.previous = chosen
        self.usage[chosen] = self.usage.get(chosen, 0) + 1
        return kind, chosen


@dataclass
class RenderResult:
    output_path: Path
    layouts: list[str]


class DeckRenderer:
    def __init__(self, deck: DeckSpec) -> None:
        self.deck = deck
        self.theme = deck.theme
        self.presentation = Presentation()
        self.presentation.slide_width = Inches(SLIDE_W)
        self.presentation.slide_height = Inches(SLIDE_H)
        self.selector = LayoutSelector()
        self.layouts: list[str] = []

    def render(self, output_path: str | Path) -> RenderResult:
        for index, spec in enumerate(self.deck.slides, start=1):
            self.theme = self.deck.theme_variants.get(spec.theme_key, self.deck.theme)
            kind, layout = self.selector.select(spec, index)
            spec.kind = kind
            self.layouts.append(layout)
            slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])
            getattr(self, f"_layout_{layout.replace('-', '_')}")(slide, spec, index)
            _add_notes(slide, spec)
        self.presentation.core_properties.title = self.deck.title
        self.presentation.core_properties.subject = self.deck.learning_objective
        self.presentation.core_properties.language = self.deck.language
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.presentation.save(str(output))
        return RenderResult(output_path=output, layouts=self.layouts)

    def _base(self, slide, spec, number, *, background=None, light=False):
        _slide_background(slide, background or self.theme.background)
        _add_header(slide, spec, self.theme, number, light=light)
        _add_footer(slide, self.deck, number, light=light, theme=self.theme)

    def _layout_cover_hero(self, slide, spec, number):
        _slide_background(slide, self.theme.ink)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 4.35, SLIDE_H, self.theme.primary)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 4.35, 0, 0.1, SLIDE_H, self.theme.accent)
        _add_text(slide, self.deck.program_name or "برنامج تدريبي", 7.9, 0.8, 4.7, 0.4, theme=self.theme, size=16, color=self.theme.accent, bold=True)
        _add_text(slide, spec.title, 5.0, 1.55, 7.55, 1.75, theme=self.theme, size=50, color=self.theme.paper, bold=True, valign=MSO_ANCHOR.MIDDLE)
        if spec.subtitle or self.deck.subtitle:
            _add_text(slide, spec.subtitle or self.deck.subtitle, 5.05, 3.5, 7.35, 0.9, theme=self.theme, size=23, color=self.theme.paper)
        _add_text(slide, self.deck.audience, 5.05, 5.85, 5.8, 0.45, theme=self.theme, size=15, color=self.theme.paper)
        _add_circle_label(slide, "01", 1.34, 2.18, 1.7, self.theme, fill=self.theme.ink, color=self.theme.paper, size=33)
        _add_footer(slide, self.deck, number, light=True, theme=self.theme)

    def _layout_cover_split(self, slide, spec, number):
        _slide_background(slide, self.theme.background)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 7.3, 0, 6.05, SLIDE_H, self.theme.primary)
        for offset in (0.6, 1.08, 1.56):
            shape = _add_shape(slide, MSO_SHAPE.OVAL, 8.3 + offset, 4.65 - offset / 3, 1.05, 1.05, self.theme.accent)
        _add_text(slide, spec.title, 0.75, 1.35, 5.95, 2.0, theme=self.theme, size=50, color=self.theme.ink, bold=True, valign=MSO_ANCHOR.MIDDLE)
        _add_text(slide, spec.subtitle or self.deck.subtitle, 0.78, 3.67, 5.55, 1.0, theme=self.theme, size=22, color=self.theme.muted)
        _add_text(slide, self.deck.program_name or "مادة تدريبية", 8.05, 1.2, 4.35, 0.5, theme=self.theme, size=18, color=self.theme.paper, bold=True)
        _add_text(slide, self.deck.learning_objective, 8.05, 2.15, 4.35, 2.1, theme=self.theme, size=25, color=self.theme.paper, bold=True, valign=MSO_ANCHOR.MIDDLE)
        _add_footer(slide, self.deck, number, theme=self.theme)

    def _layout_section_number(self, slide, spec, number):
        _slide_background(slide, self.theme.primary)
        _add_text(slide, f"{number:02d}", 0.75, 0.6, 3.5, 2.0, theme=self.theme, size=92, color=self.theme.accent, bold=True, align=PP_ALIGN.LEFT, rtl=False)
        _add_text(slide, spec.title, 4.5, 1.55, 7.9, 1.5, theme=self.theme, size=47, color=self.theme.paper, bold=True, valign=MSO_ANCHOR.MIDDLE)
        _add_text(slide, spec.subtitle or spec.body, 4.52, 3.25, 7.1, 1.25, theme=self.theme, size=24, color=self.theme.paper)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 4.52, 5.3, 2.5, 0.08, self.theme.accent)
        _add_footer(slide, self.deck, number, light=True, theme=self.theme)

    def _layout_section_band(self, slide, spec, number):
        _slide_background(slide, self.theme.background)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 2.02, SLIDE_W, 3.25, self.theme.ink)
        _add_text(slide, spec.eyebrow or "وحدة تدريبية", 8.0, 1.25, 4.6, 0.4, theme=self.theme, size=15, color=self.theme.primary, bold=True)
        _add_text(slide, spec.title, 2.0, 2.48, 9.35, 1.25, theme=self.theme, size=45, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        _add_text(slide, spec.subtitle or spec.body, 3.0, 4.0, 7.35, 0.65, theme=self.theme, size=21, color=self.theme.paper, align=PP_ALIGN.CENTER)
        _add_footer(slide, self.deck, number, theme=self.theme)

    def _layout_statement_center(self, slide, spec, number):
        self._base(slide, spec, number)
        _add_text(slide, "“", 10.95, 1.53, 1.0, 0.9, theme=self.theme, size=58, color=self.theme.accent, bold=True)
        statement = spec.body or spec.takeaway or spec.subtitle
        _add_text(slide, statement, 1.55, 2.15, 10.25, 2.45, theme=self.theme, size=37, color=self.theme.ink, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 5.34, 5.15, 2.65, 0.09, self.theme.primary)

    def _layout_statement_sidebar(self, slide, spec, number):
        self._base(slide, spec, number)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.7, 1.72, 3.85, 4.62, self.theme.primary)
        _add_circle_label(slide, "!", 10.05, 2.42, 1.25, self.theme, fill=self.theme.accent, color=self.theme.ink, size=35)
        _add_text(slide, spec.takeaway or "الفكرة المحورية", 9.2, 4.1, 2.85, 1.25, theme=self.theme, size=26, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        _add_text(slide, spec.body or spec.subtitle, 0.85, 2.02, 7.15, 2.75, theme=self.theme, size=34, color=self.theme.ink, bold=True, valign=MSO_ANCHOR.MIDDLE)
        if spec.bullets:
            _add_multiline(slide, spec.bullets[:3], 0.92, 5.05, 7.0, 1.25, theme=self.theme, size=18, color=self.theme.muted, bullet=True)

    def _layout_concept_cards(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:4]
        count = len(items)
        gap = 0.24
        card_w = (11.9 - gap * (count - 1)) / count
        for visual_index, item in enumerate(reversed(items)):
            x = 0.72 + visual_index * (card_w + gap)
            card = _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.86, card_w, 4.72, self.theme.paper, self.theme.background)
            _set_alt_text(card, item.title or item.label)
            _add_circle_label(slide, item.label or str(count - visual_index), x + card_w - 0.72, 2.18, 0.48, self.theme, size=13)
            _add_text(slide, item.value, x + 0.22, 2.12, card_w - 1.05, 0.58, theme=self.theme, size=27, color=self.theme.accent, bold=True, align=PP_ALIGN.LEFT, rtl=False)
            _add_text(slide, item.title or item.label, x + 0.22, 2.98, card_w - 0.44, 0.9, theme=self.theme, size=23, color=self.theme.ink, bold=True, valign=MSO_ANCHOR.MIDDLE)
            _add_text(slide, item.text or item.note, x + 0.22, 4.02, card_w - 0.44, 1.75, theme=self.theme, size=17.5, color=self.theme.muted)
        if spec.takeaway:
            _add_text(slide, spec.takeaway, 1.1, 6.58, 11.1, 0.32, theme=self.theme, size=14, color=self.theme.primary, bold=True, align=PP_ALIGN.CENTER)

    def _layout_concept_split(self, slide, spec, number):
        self._base(slide, spec, number)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 7.35, 1.78, 5.15, 4.95, self.theme.ink)
        _add_text(slide, spec.body or spec.takeaway, 7.78, 2.25, 4.28, 3.55, theme=self.theme, size=29, color=self.theme.paper, bold=True, valign=MSO_ANCHOR.MIDDLE)
        lines = spec.bullets or [_item_text(item) for item in spec.items]
        _add_multiline(slide, lines[:5], 0.85, 2.02, 5.85, 3.9, theme=self.theme, size=21, color=self.theme.ink, bullet=True, spacing=14)
        _add_shape(slide, MSO_SHAPE.RECTANGLE, 0.95, 6.15, 5.65, 0.1, self.theme.accent)

    def _layout_concept_focus(self, slide, spec, number):
        self._base(slide, spec, number)
        _add_shape(slide, MSO_SHAPE.OVAL, 4.75, 2.05, 3.8, 3.8, self.theme.primary)
        _add_text(slide, spec.takeaway or spec.body or spec.title, 5.2, 2.73, 2.9, 2.1, theme=self.theme, size=27, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        items = _fit_items(spec.items, spec.bullets)[:4]
        positions = [(9.2, 1.95), (9.15, 4.65), (0.85, 1.95), (0.9, 4.65)]
        for item, (x, y) in zip(items, positions):
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.2, 1.28, self.theme.paper, self.theme.background)
            _add_text(slide, item.title or item.label, x + 0.18, y + 0.12, 2.85, 0.4, theme=self.theme, size=18, color=self.theme.ink, bold=True)
            _add_text(slide, item.text or item.note, x + 0.18, y + 0.58, 2.85, 0.52, theme=self.theme, size=14.5, color=self.theme.muted)

    def _layout_process_horizontal(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:5]
        count = len(items)
        xs = [1.0 + i * (11.2 / max(count - 1, 1)) for i in range(count)]
        y = 3.2
        if count > 1:
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xs[0]), Inches(y + 0.42), Inches(xs[-1]), Inches(y + 0.42))
            _line(connector, self.theme.primary, 3.2)
        for idx, (item, x) in enumerate(zip(items, xs), start=1):
            _add_circle_label(slide, str(idx), x - 0.42, y, 0.84, self.theme, fill=self.theme.primary, size=21)
            _add_text(slide, item.title or item.label, x - 1.02, y + 1.1, 2.05, 0.65, theme=self.theme, size=18, color=self.theme.ink, bold=True, align=PP_ALIGN.CENTER)
            _add_text(slide, item.text or item.note, x - 1.12, y + 1.82, 2.25, 0.85, theme=self.theme, size=14, color=self.theme.muted, align=PP_ALIGN.CENTER)

    def _layout_process_stair(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:5]
        card_w = 2.18
        for idx, item in enumerate(reversed(items)):
            visual = len(items) - 1 - idx
            x = 0.78 + idx * 2.43
            y = 4.75 - visual * 0.58
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, 1.58, self.theme.paper, self.theme.background)
            _add_text(slide, str(visual + 1), x + 0.15, y + 0.12, 0.35, 0.35, theme=self.theme, size=15, color=self.theme.accent, bold=True, align=PP_ALIGN.LEFT, rtl=False)
            _add_text(slide, item.title or item.label, x + 0.15, y + 0.48, 1.88, 0.44, theme=self.theme, size=17.5, color=self.theme.ink, bold=True)
            _add_text(slide, item.text or item.note, x + 0.15, y + 0.94, 1.88, 0.48, theme=self.theme, size=12.8, color=self.theme.muted)

    def _layout_timeline_ribbon(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:6]
        count = len(items)
        width = 11.8 / count
        for idx, item in enumerate(reversed(items)):
            x = 0.75 + idx * width
            color = self.theme.primary if idx % 2 == 0 else self.theme.ink
            _add_shape(slide, MSO_SHAPE.CHEVRON, x, 2.35, width + 0.12, 1.05, color)
            _add_text(slide, item.label or str(count - idx), x + 0.15, 2.59, width - 0.3, 0.42, theme=self.theme, size=16, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER)
            _add_text(slide, item.title or item.value, x + 0.1, 3.72, width - 0.2, 0.62, theme=self.theme, size=17, color=self.theme.ink, bold=True, align=PP_ALIGN.CENTER)
            _add_text(slide, item.text or item.note, x + 0.1, 4.45, width - 0.2, 1.0, theme=self.theme, size=13.5, color=self.theme.muted, align=PP_ALIGN.CENTER)

    def _layout_timeline_lanes(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:5]
        for idx, item in enumerate(items):
            y = 1.72 + idx * 0.94
            x = 1.0 + (idx % 2) * 0.45
            _add_circle_label(slide, item.label or str(idx + 1), 11.55, y + 0.08, 0.55, self.theme, size=13)
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(11.55), Inches(y + 0.36), Inches(x + 0.18), Inches(y + 0.36))
            _line(line, self.theme.background if idx % 2 else self.theme.accent, 2.3)
            _add_text(slide, item.title or item.value, 8.0, y, 3.25, 0.42, theme=self.theme, size=18, color=self.theme.ink, bold=True)
            _add_text(slide, item.text or item.note, x, y, 6.5, 0.58, theme=self.theme, size=15.2, color=self.theme.muted)

    def _layout_comparison_columns(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:3]
        count = len(items)
        gap = 0.25
        width = (11.8 - gap * (count - 1)) / count
        for idx, item in enumerate(reversed(items)):
            x = 0.75 + idx * (width + gap)
            fill = self.theme.primary if idx == 0 else self.theme.paper
            text_color = self.theme.paper if idx == 0 else self.theme.ink
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.8, width, 4.82, fill, self.theme.background if idx else None)
            _add_text(slide, item.label or item.title, x + 0.24, 2.13, width - 0.48, 0.65, theme=self.theme, size=24, color=text_color, bold=True, align=PP_ALIGN.CENTER)
            _add_text(slide, item.value, x + 0.24, 2.95, width - 0.48, 0.68, theme=self.theme, size=30, color=self.theme.accent, bold=True, align=PP_ALIGN.CENTER, rtl=False)
            _add_text(slide, item.text or item.note, x + 0.36, 3.88, width - 0.72, 1.95, theme=self.theme, size=18, color=text_color, align=PP_ALIGN.CENTER)

    def _layout_comparison_balance(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:2]
        while len(items) < 2:
            items.append(ContentItem(title="الخيار الآخر"))
        for idx, item in enumerate(items):
            x = 7.35 if idx == 0 else 0.8
            fill = self.theme.primary if idx == 0 else self.theme.ink
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.02, 5.18, 3.95, fill)
            _add_text(slide, item.title or item.label, x + 0.35, 2.4, 4.48, 0.7, theme=self.theme, size=27, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER)
            _add_text(slide, item.text or item.note, x + 0.5, 3.35, 4.18, 1.75, theme=self.theme, size=19, color=self.theme.paper, align=PP_ALIGN.CENTER)
        _add_circle_label(slide, "↔", 6.08, 3.24, 1.16, self.theme, fill=self.theme.accent, color=self.theme.ink, size=28)

    def _layout_metrics_grid(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:6]
        cols = 3
        rows = math.ceil(len(items) / cols)
        for idx, item in enumerate(items):
            row, col = divmod(idx, cols)
            visual_col = cols - 1 - col
            x = 0.8 + visual_col * 4.02
            y = 1.82 + row * 2.25
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.7, 1.93, self.theme.paper, self.theme.background)
            _add_text(slide, item.value or item.label, x + 0.25, y + 0.17, 3.2, 0.72, theme=self.theme, size=32, color=self.theme.primary, bold=True, align=PP_ALIGN.CENTER, rtl=False)
            _add_text(slide, item.title or item.text, x + 0.28, y + 1.02, 3.14, 0.56, theme=self.theme, size=17, color=self.theme.ink, bold=True, align=PP_ALIGN.CENTER)

    def _layout_metrics_hero(self, slide, spec, number):
        self._base(slide, spec, number, background=self.theme.ink, light=True)
        items = _fit_items(spec.items, spec.bullets)[:4]
        hero = items[0]
        _add_text(slide, hero.value or hero.label, 6.75, 1.72, 5.55, 1.65, theme=self.theme, size=64, color=self.theme.accent, bold=True, align=PP_ALIGN.CENTER, rtl=False, valign=MSO_ANCHOR.MIDDLE)
        _add_text(slide, hero.title or hero.text, 7.05, 3.4, 5.0, 1.05, theme=self.theme, size=25, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER)
        for idx, item in enumerate(items[1:4]):
            y = 1.95 + idx * 1.42
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, y, 5.0, 1.05, self.theme.primary)
            _add_text(slide, item.value or item.label, 1.12, y + 0.18, 1.45, 0.55, theme=self.theme, size=24, color=self.theme.accent, bold=True, align=PP_ALIGN.LEFT, rtl=False)
            _add_text(slide, item.title or item.text, 2.72, y + 0.16, 2.82, 0.6, theme=self.theme, size=17.5, color=self.theme.paper, bold=True)

    def _layout_case_decision(self, slide, spec, number):
        self._base(slide, spec, number)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 7.65, 1.78, 4.85, 4.9, self.theme.ink)
        _add_text(slide, "الموقف", 10.78, 2.05, 1.2, 0.4, theme=self.theme, size=16, color=self.theme.accent, bold=True)
        _add_text(slide, spec.body or spec.subtitle, 8.05, 2.65, 4.05, 2.7, theme=self.theme, size=24, color=self.theme.paper, bold=True, valign=MSO_ANCHOR.MIDDLE)
        items = _fit_items(spec.items, spec.bullets)[:3]
        for idx, item in enumerate(items):
            y = 1.88 + idx * 1.43
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, y, 5.95, 1.12, self.theme.paper, self.theme.background)
            _add_circle_label(slide, str(idx + 1), 6.04, y + 0.24, 0.48, self.theme, size=13)
            _add_text(slide, item.title or item.label, 3.7, y + 0.15, 2.1, 0.36, theme=self.theme, size=17.5, color=self.theme.ink, bold=True)
            _add_text(slide, item.text or item.note, 1.08, y + 0.57, 4.72, 0.35, theme=self.theme, size=14.5, color=self.theme.muted)
        _add_text(slide, spec.takeaway or "ما القرار؟ وما المفاضلة التي تبرره؟", 0.88, 6.25, 5.9, 0.42, theme=self.theme, size=17, color=self.theme.primary, bold=True)

    def _layout_case_scenario(self, slide, spec, number):
        self._base(slide, spec, number)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 1.82, 11.65, 2.05, self.theme.paper, self.theme.background)
        _add_text(slide, spec.body or spec.subtitle, 1.25, 2.18, 10.85, 1.3, theme=self.theme, size=26, color=self.theme.ink, bold=True, valign=MSO_ANCHOR.MIDDLE)
        items = _fit_items(spec.items, spec.bullets)[:3]
        labels = ["الإشارة", "القرار", "المعيار"]
        for idx, item in enumerate(reversed(items)):
            x = 0.92 + idx * 3.92
            _add_text(slide, labels[min(idx, 2)], x, 4.28, 3.45, 0.36, theme=self.theme, size=14, color=self.theme.primary, bold=True, align=PP_ALIGN.CENTER)
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.73, 3.45, 1.55, self.theme.ink if idx == 1 else self.theme.primary)
            _add_text(slide, item.title or item.text, x + 0.25, 4.98, 2.95, 0.95, theme=self.theme, size=19, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

    def _layout_activity_workshop(self, slide, spec, number):
        self._base(slide, spec, number)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.9, 1.75, 2.55, 1.2, self.theme.accent)
        _add_text(slide, f"{spec.duration_minutes or 10}\nدقائق", 10.2, 1.9, 1.95, 0.82, theme=self.theme, size=22, color=self.theme.ink, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        _add_text(slide, spec.body or spec.subtitle, 0.92, 1.92, 8.35, 1.15, theme=self.theme, size=27, color=self.theme.ink, bold=True, valign=MSO_ANCHOR.MIDDLE)
        items = _fit_items(spec.items, spec.bullets)[:4]
        for idx, item in enumerate(items):
            y = 3.35 + idx * 0.77
            _add_circle_label(slide, str(idx + 1), 11.7, y, 0.5, self.theme, size=13)
            _add_text(slide, item.title or item.text, 5.95, y + 0.04, 5.4, 0.42, theme=self.theme, size=18.5, color=self.theme.ink, bold=True)
        _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.25, 4.35, 1.1, self.theme.primary)
        _add_text(slide, spec.takeaway or "المخرج: قرار مبرر وقابل للمراجعة", 1.2, 5.5, 3.75, 0.55, theme=self.theme, size=18, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER)

    def _layout_activity_canvas(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:4]
        quadrants = [(6.8, 1.8), (0.85, 1.8), (6.8, 4.32), (0.85, 4.32)]
        for idx, ((x, y), item) in enumerate(zip(quadrants, items)):
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 5.55, 2.05, self.theme.paper, self.theme.background)
            _add_text(slide, item.label or str(idx + 1), x + 4.75, y + 0.22, 0.5, 0.4, theme=self.theme, size=15, color=self.theme.primary, bold=True)
            _add_text(slide, item.title or item.text, x + 0.35, y + 0.67, 4.75, 0.78, theme=self.theme, size=21, color=self.theme.ink, bold=True, valign=MSO_ANCHOR.MIDDLE)

    def _layout_quiz_options(self, slide, spec, number):
        self._base(slide, spec, number, background=self.theme.ink, light=True)
        _add_text(slide, spec.body or spec.subtitle or spec.title, 1.1, 1.65, 11.1, 1.28, theme=self.theme, size=31, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        choices = spec.choices or spec.bullets[:4]
        labels = "ABCD"
        for idx, choice in enumerate(choices[:4]):
            row, col = divmod(idx, 2)
            visual_col = 1 - col
            x = 0.95 + visual_col * 6.1
            y = 3.22 + row * 1.38
            fill = self.theme.primary if idx % 2 == 0 else "16324A"
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 5.55, 1.05, fill)
            _add_circle_label(slide, labels[idx], x + 4.72, y + 0.18, 0.56, self.theme, fill=self.theme.accent, color=self.theme.ink, size=14)
            _add_text(slide, choice, x + 0.35, y + 0.2, 4.15, 0.58, theme=self.theme, size=18.5, color=self.theme.paper, bold=True, valign=MSO_ANCHOR.MIDDLE)

    def _layout_tool_canvas(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:5]
        widths = [2.15, 2.15, 2.15, 2.15, 2.15]
        for idx, item in enumerate(reversed(items)):
            x = 0.8 + idx * 2.42
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.05, widths[idx], 3.95, self.theme.paper, self.theme.background)
            _add_text(slide, item.label or str(len(items) - idx), x + 0.18, 2.25, 0.45, 0.38, theme=self.theme, size=14, color=self.theme.accent, bold=True, align=PP_ALIGN.LEFT, rtl=False)
            _add_text(slide, item.title, x + 0.18, 2.8, 1.79, 0.75, theme=self.theme, size=19, color=self.theme.ink, bold=True, align=PP_ALIGN.CENTER)
            _add_text(slide, item.text or item.note, x + 0.18, 3.78, 1.79, 1.35, theme=self.theme, size=14.5, color=self.theme.muted, align=PP_ALIGN.CENTER)
        if spec.takeaway:
            _add_text(slide, spec.takeaway, 1.05, 6.3, 11.2, 0.45, theme=self.theme, size=16, color=self.theme.primary, bold=True, align=PP_ALIGN.CENTER)

    def _layout_tool_steps(self, slide, spec, number):
        self._base(slide, spec, number)
        items = _fit_items(spec.items, spec.bullets)[:4]
        for idx, item in enumerate(items):
            y = 1.8 + idx * 1.18
            fill = self.theme.primary if idx == 0 else self.theme.paper
            text_color = self.theme.paper if idx == 0 else self.theme.ink
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.05 + idx * 0.28, y, 11.2 - idx * 0.56, 0.92, fill, None if idx == 0 else self.theme.background)
            _add_text(slide, item.title or item.label, 7.4, y + 0.2, 4.0 - idx * 0.15, 0.42, theme=self.theme, size=18, color=text_color, bold=True)
            _add_text(slide, item.text or item.note, 1.42 + idx * 0.28, y + 0.2, 5.55, 0.42, theme=self.theme, size=15.5, color=text_color)

    def _layout_summary_checklist(self, slide, spec, number):
        self._base(slide, spec, number)
        lines = spec.bullets or [_item_text(item) for item in spec.items]
        for idx, line in enumerate(lines[:6]):
            y = 1.78 + idx * 0.78
            _add_circle_label(slide, "✓", 11.63, y, 0.48, self.theme, fill=self.theme.success, size=13)
            _add_text(slide, line, 2.0, y + 0.04, 9.25, 0.42, theme=self.theme, size=20, color=self.theme.ink, bold=True)
        if spec.takeaway:
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.15, 6.35, 11.0, 0.55, self.theme.primary)
            _add_text(slide, spec.takeaway, 1.4, 6.43, 10.5, 0.34, theme=self.theme, size=15.5, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER)

    def _layout_summary_commitment(self, slide, spec, number):
        self._base(slide, spec, number, background=self.theme.primary, light=True)
        _add_text(slide, spec.takeaway or spec.body or "ما الخطوة التي ستبدأ بها؟", 1.15, 1.83, 11.05, 1.35, theme=self.theme, size=39, color=self.theme.paper, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        items = _fit_items(spec.items, spec.bullets)[:3]
        for idx, item in enumerate(reversed(items)):
            x = 1.02 + idx * 3.93
            _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.72, 3.55, 1.75, self.theme.paper)
            _add_text(slide, item.label or str(len(items) - idx), x + 2.78, 3.98, 0.42, 0.35, theme=self.theme, size=14, color=self.theme.accent, bold=True)
            _add_text(slide, item.title or item.text, x + 0.28, 4.38, 2.99, 0.62, theme=self.theme, size=19, color=self.theme.ink, bold=True, align=PP_ALIGN.CENTER)


def render_deck(deck: DeckSpec, output_path: str | Path) -> RenderResult:
    return DeckRenderer(deck).render(output_path)
